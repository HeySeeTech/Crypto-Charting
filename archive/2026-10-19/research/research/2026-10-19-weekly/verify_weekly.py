import json,shutil
from pathlib import Path
from openpyxl import load_workbook
root=Path('/Users/jss/Desktop/workspace/XCreator/crypto-charting-data')
out=root/'weekly/2026-10-19'
checks=[]
for s in ['UNI','ZEC']:
 j=json.loads((out/(s+'.json')).read_text())
 assert j['report_period_utc']=={'start':'2026-10-12','end':'2026-10-18','observations':7,'basis':'completed UTC days; live quote is separately timestamped'}
 assert j['collected_at'].startswith('2026-10-19T11:59:')
 assert j['weekly_comparison']['status']=='unavailable'
 w=load_workbook(out/(s+'.xlsx'),data_only=True)
 errors=[(ws.title,c.coordinate,c.value) for ws in w for row in ws for c in row if c.data_type=='e']
 assert not errors,errors
 assert w.worksheets[0]['A10'].value==j['metrics']['price_usd']['value']
 f=load_workbook(out/(s+'.xlsx'),data_only=False)
 charts=sum(len(ws._charts) for ws in f)
 assert charts>=2
 if s=='ZEC':
  assert abs(w.worksheets[0]['D16'].value-(-37.9))<1e-6
  assert w.worksheets[0]['G16'].value==547192.63
 checks.append({'symbol':s,'checks':'JSON period/provider timestamp, cached formulas, chart presence, source reconciliation','charts':charts,'errors':errors,'status':'pass'})
for s in ['UNI','ZEC']:
 shutil.copy2(out/(s+'.xlsx'),root/'latest'/(s+'.xlsx'))
 shutil.copy2(root/'2026-10-19'/(s+'.json'),root/'latest'/(s+'.json'))
(out/'validation.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(checks,ensure_ascii=False))
